import fitz
from docx import Document
from app.ingestion import extract, preview
from app.research import html_document, allowed


def test_pdf_pages_preview_scan_and_duplicate_identity(tmp_path):
    path = tmp_path / "test.pdf"
    with fitz.open() as d:
        d.new_page().insert_text((30, 30), "Scoring rules: grade weight 70%, bonus weight 30%.")
        d.new_page()
        d.save(path)
    doc = extract(path)
    assert len(doc["units"]) == 2
    assert doc["units"][1]["needs_ocr"]
    assert preview(doc, 1).startswith(b"\x89PNG")
    assert extract(path, "different-name.pdf")["id"] == doc["id"]


def test_docx_keeps_paragraph_and_table_locations(tmp_path):
    path = tmp_path / "rule.docx"
    d = Document()
    d.add_paragraph("First rule")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Level"
    table.cell(0, 1).text = "Points"
    d.add_paragraph("Final rule")
    d.save(path)
    out = extract(path)
    assert [x["text"] for x in out["units"]] == ["First rule", "Level | Points", "Final rule"]
    assert all("页" not in u["label"] for u in out["units"])


def test_official_attachment_resolution_and_domains():
    raw = b'<html><title>Contest</title><article>Registration ends 2026-10-20.<a href="../files/rules.pdf">Rules</a><a href="https://evil.test/file.docx">External</a></article></html>'
    doc, links = html_document(raw, "https://school.edu.cn/news/notice.html")
    assert links[0]["url"] == "https://school.edu.cn/files/rules.pdf"
    assert allowed(links[0]["url"], ["edu.cn"])
    assert not allowed("https://school.edu.cn.evil.test/rules.pdf", ["school.edu.cn"])
    assert "Registration ends" in doc["units"][0]["text"]
